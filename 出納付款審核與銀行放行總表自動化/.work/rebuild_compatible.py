from pathlib import Path
from copy import copy
from collections import OrderedDict
from decimal import Decimal
import openpyxl
from datetime import datetime
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import RowBreak,ColBreak
root=Path.cwd()
erp=openpyxl.load_workbook(root/'素材_應付明細_台幣.xlsx',read_only=True,data_only=True)
groups=OrderedDict()
for row in erp.active.iter_rows(min_row=2,values_only=True):
 due=row[7]
 if isinstance(due,datetime):due=due.strftime('%Y/%m/%d')
 if str(due).strip()!='2026/09/30':continue
 key=(str(row[4]).strip(),str(row[2]).strip())
 if key not in groups:groups[key]=[0,Decimal(0)]
 groups[key][0]+=1;groups[key][1]+=Decimal(str(row[13]))
w=openpyxl.load_workbook(root/'素材_放行總表.xlsx',keep_links=False)
s=w.active
for name in list(w.defined_names):
 if name=='銀行名稱':del w.defined_names[name]
for r in range(4,9):
 for c in range(1,11):
  dest=s.cell(r,c);dest._style=copy(s.cell(4,c)._style);dest.value=None
for i,((voucher,vendor),(count,amount)) in enumerate(groups.items(),4):
 values=[i-3,voucher,datetime(2026,9,30),f'應付 {vendor}（{count}筆發票）',0,float(amount),0,f'=SUM(F{i}:G{i})',None,None]
 for c,val in enumerate(values,1):s.cell(i,c,val)
 for c in range(5,10):s.cell(i,c).number_format='#,##0;[Red](#,##0);-'
 s.row_dimensions[i].height=45
s['A2']='臺企銀 臺幣單筆付款 放行總表'
s['J2']='FIN115-175'
for c in ['E','F','G','H']:s[f'{c}9']=f'=SUM({c}4:{c}8)'
s['I9']=None
for row in s.iter_rows(min_row=10,max_row=10):
 for cell in row:cell.value=None
s.merge_cells('A10:J10');s['A10']='本期手續費由公司負擔；單筆大額款項請主管覆核放行。'
s['A10'].font=copy(s['D4'].font);s['A10'].alignment=openpyxl.styles.Alignment(horizontal='left',vertical='center')
s.row_dimensions[10].height=32
s['A11']='核准：                         覆核：                         審核：                         經辦：'
s.row_dimensions[11].height=36
s.print_area='A1:J12';s.print_title_rows='1:3'
s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4
s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1;s.sheet_properties.pageSetUpPr.fitToPage=True
s.row_breaks=RowBreak();s.col_breaks=ColBreak()
out=root/'outputs/01a10bc7-8ae6-7fc0-82fa-7915b40e0642/出納付款放行總表_已完成_Excel相容版.xlsx'
w.save(out)
check=openpyxl.load_workbook(out,data_only=False,keep_links=False)
assert sum(check.active.cell(r,6).value for r in range(4,9))==4040257
assert check.active['H9'].value=='=SUM(H4:H8)'
assert check.active['A10'].value.startswith('本期手續費')
assert len(groups)==5 and sum(x[0] for x in groups.values())==10
print(out)
