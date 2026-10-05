from pathlib import Path
from copy import copy
from collections import OrderedDict
from decimal import Decimal
from datetime import datetime,date
import openpyxl
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.styles import Alignment
from openpyxl.workbook.properties import CalcProperties
root=Path.cwd()
erp=openpyxl.load_workbook(root/'素材_應付明細_台幣.xlsx',read_only=True,data_only=True)
groups=OrderedDict()
for row in erp.active.iter_rows(min_row=2,values_only=True):
 due=row[7].date() if isinstance(row[7],datetime) else row[7] if isinstance(row[7],date) else datetime.strptime(str(row[7]).strip(),'%Y/%m/%d').date()
 if not date(2026,10,1)<=due<=date(2026,10,6):continue
 key=(str(row[4]).strip(),str(row[2]).strip())
 groups.setdefault(key,[0,Decimal(0)])
 groups[key][0]+=1;groups[key][1]+=Decimal(str(row[13]))
assert len(groups)==7 and sum(x[0] for x in groups.values())==8 and sum(x[1] for x in groups.values())==1644295
template=openpyxl.load_workbook(root/'放行總表_佔位字符樣版.xlsx',keep_links=False)
t=template.active
# Rebuild through the standard Excel serializer, preserving template styling.
w=openpyxl.Workbook();s=w.active;s.title=t.title
for col,dim in t.column_dimensions.items():
 s.column_dimensions[col]=ColumnDimension(s,index=col,width=dim.width,min=dim.min,max=dim.max,hidden=dim.hidden,outlineLevel=dim.outlineLevel,collapsed=dim.collapsed,bestFit=dim.bestFit)
s.sheet_format=copy(t.sheet_format);s.sheet_properties=copy(t.sheet_properties)
s.page_margins=copy(t.page_margins);s.print_options=copy(t.print_options)
s.oddHeader=copy(t.oddHeader);s.oddFooter=copy(t.oddFooter)
def rowstyle(dest,src):
 s.row_dimensions[dest].height=t.row_dimensions[src].height
 for c in range(1,11):
  cell=s.cell(dest,c);ref=t.cell(src,c)
  for prop in ['font','fill','border','alignment','protection']:setattr(cell,prop,copy(getattr(ref,prop)))
  cell.number_format=ref.number_format
for r in range(1,4):
 rowstyle(r,r)
 for c in range(1,11):s.cell(r,c).value=t.cell(r,c).value
s.merge_cells('B1:J1');s.merge_cells('A2:I2')
s['A2']='臺企銀 臺幣單筆付款 放行總表';s['J2']='FIN115-175'
for r,((voucher,vendor),(count,amount)) in enumerate(groups.items(),4):
 rowstyle(r,4)
 for c,val in enumerate([r-3,voucher,datetime(2026,10,6),f'應付 {vendor}（{count}筆發票）',0,int(amount),0,f'=SUM(F{r}:G{r})',None,None],1):s.cell(r,c).value=val
 s.row_dimensions[r].height=45
 s.cell(r,3).number_format='m/d'
end=3+len(groups);total=end+1;note=total+1;sign=total+2
rowstyle(total,9);s.merge_cells(start_row=total,start_column=1,end_row=total,end_column=4)
s.cell(total,1).value='合    計：'
for col in ['E','F','G','H']:s[f'{col}{total}']=f'=SUM({col}4:{col}{end})'
rowstyle(note,10);s.merge_cells(start_row=note,start_column=1,end_row=note,end_column=10)
s.cell(note,1).value='本期手續費由公司負擔；單筆大額款項請主管覆核放行。'
s.cell(note,1).font=copy(t['D4'].font);s.cell(note,1).alignment=Alignment(horizontal='left',vertical='center')
s.row_dimensions[note].height=32
rowstyle(sign,11);s.merge_cells(start_row=sign,start_column=1,end_row=sign,end_column=10)
s.cell(sign,1).value='核准：                         覆核：                         審核：                         經辦：'
s.row_dimensions[sign].height=36
s.print_area=f'A1:J{sign+1}';s.print_title_rows='1:3'
s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
s.sheet_properties.pageSetUpPr.fitToPage=True
w.calculation=CalcProperties(fullCalcOnLoad=True)
out=root/'outputs/01a10bc7-8ae6-7fc0-82fa-7915b40e0642/出納付款放行總表_已完成_20261001_20261006.xlsx'
w.save(out)
c=openpyxl.load_workbook(out,data_only=False).active
assert sum(c.cell(r,6).value for r in range(4,end+1))==1644295
assert c[f'H{total}'].value==f'=SUM(H4:H{end})'
assert c.cell(sign,1).value.startswith('核准：')
assert not any('{{' in str(cell.value) for row in c for cell in row)
print(f'完成：7 筆／8 張／1,644,295 元；資料第 4–{end} 列，合計 {total}，備註 {note}，簽章 {sign}。')
